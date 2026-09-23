"""Рендер технического задания: Word для человека, Markdown для чата и графа.

Два выхода из одного `TechnicalSpec`, потому что документ живёт в двух местах.
В .docx его открывают, печатают и подписывают; .md показывает то же самое там,
где Word не откроешь — в ленте чата и в карточке «Постановка».

Оформление намеренно скромное: ГОСТ 19.201-78 задаёт СОДЕРЖАНИЕ разделов, а
требования к оформлению листов лежат в ГОСТ 19.106-78, который мы не заявляем и
не проверяем. Обещать в колонтитуле соответствие тому, чего никто не сверял, —
хуже, чем не обещать ничего. Но то немногое, что оформление обещает, оно должно
держать: один шрифт, чёрный текст, поля под подшивку. Шаблон Word по умолчанию
не держит ничего из этого — заголовки в нём синие и набраны другой гарнитурой,
и документ читается как веб-страница.

Исходный запрос заказчика в документ НЕ переносится. Он неформальный («собери
данные», «автоматизируй…»), и в официальном тексте выглядит тем, чем является, —
чужой репликой. Его место — рамка исследования и граф, где он и хранится; в ТЗ
попадает то, что из него сформулировано.
"""
from __future__ import annotations

import io
from typing import List, Optional, Tuple

from CoScientist.context_init.tz import NOT_SET, TZSection, TechnicalSpec

#: Что печатается на титуле над темой.
_STANDARD = "ГОСТ 19.201-78"

#: Гарнитура и размеры. Times New Roman 12 — регистр, в котором такие документы
#: читают; python-docx ставит Calibri 11, и ТЗ выглядит как записка.
_FONT = "Times New Roman"
_BODY_PT = 12
_HEADING_PT = {1: 14, 2: 13, 3: 12, 4: 12}


def _visible(section: TZSection) -> bool:
    """Печатать ли раздел. Печатаются все — включая пустые.

    Пустой раздел в ТЗ не мусор, а вопрос заказчику: «основание для работы не
    задано» сообщает читателю ровно то, что произошло. Молча выброшенный раздел
    сообщил бы, что его не требовалось.
    """
    return bool(section.number)


# ── Markdown ────────────────────────────────────────────────────────────────

def _md_section(section: TZSection, level: int = 2) -> List[str]:
    out = ["", "#" * level + f" {section.number}. {section.title}"]
    if section.body:
        out += ["", section.body]
    if section.rows:
        out += ["", "| | |", "|---|---|"]
        out += [f"| {k} | {v} |" for k, v in section.rows]
    for sub in section.subsections:
        out += _md_section(sub, level + 1)
    return out


def render_tz_markdown(spec: TechnicalSpec) -> str:
    """Тот же документ, что и .docx, в виде Markdown."""
    lines = [
        "# Техническое задание",
        "",
        f"**{spec.topic}**",
        "",
        f"Составлено по {_STANDARD}.",
        "",
        "| | |",
        "|---|---|",
        f"| Заказчик | {spec.customer} |",
        f"| Основание | {spec.basis} |",
    ]
    for section in spec.sections:
        if _visible(section):
            lines += _md_section(section)
    missing = spec.unfilled()
    if missing:
        lines += ["", "---", "",
                  "Разделы, для которых сведения не заданы: "
                  + ", ".join(missing) + ". "
                  "Они заполняются оператором в форме «Основание и приёмка» "
                  "рамки исследования."]
    return "\n".join(lines) + "\n"


# ── Word ────────────────────────────────────────────────────────────────────

def _force_font(style, size_pt: int, bold: bool) -> None:
    """Прибить гарнитуру и чёрный цвет к стилю, поверх темы документа.

    Через XML, а не через `style.font`: шаблон Word задаёт заголовкам шрифт и
    цвет ссылками на тему (`w:asciiTheme`, `w:themeColor="accent1"`), и Word
    предпочитает ссылку явному значению. Пока ссылка на месте, заголовок
    остаётся синим, сколько бы раз ему ни назначили чёрный.
    """
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    style.font.name = _FONT
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    # Курсив шаблон даёт заголовкам нижних уровней; в ТЗ заголовки прямые.
    style.font.italic = False
    style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is not None:
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attr}"), None)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), _FONT)
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attr in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn(f"w:{attr}"), None)
        color.set(qn("w:val"), "000000")


def _typography(document) -> None:
    """Шрифт, цвет, абзацный отступ и поля — один раз на весь документ."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    styles = document.styles
    normal = styles["Normal"]
    _force_font(normal, _BODY_PT, bold=False)
    body = normal.paragraph_format
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.first_line_indent = Cm(1.25)
    body.space_after = Pt(6)

    for level, size in _HEADING_PT.items():
        try:
            style = styles[f"Heading {level}"]
        except KeyError:  # шаблон без этого уровня — заголовок просто наследует
            continue
        _force_font(style, size, bold=True)
        fmt = style.paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.first_line_indent = Cm(0)
        fmt.space_before = Pt(12)
        fmt.space_after = Pt(6)
        fmt.keep_with_next = True

    # A4 и поля под подшивку слева: документ печатают и подшивают, а шаблон
    # python-docx свёрстан под Letter (21,59 × 27,94 см) — американский формат,
    # который в российской организации вылезет за поля при печати.
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)


def _plain(document, text: str = "", *, align=None, indent_cm: float = 0.0,
           bold: bool = False, italic: bool = False, size_pt: int = 0):
    """Абзац без абзацного отступа — титул, подписи, служебные строки."""
    from docx.shared import Cm, Pt

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(indent_cm)
    if align is not None:
        paragraph.alignment = align
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        if size_pt:
            run.font.size = Pt(size_pt)
    return paragraph


def _add_rows(document, rows: List[Tuple[str, str]]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)
        for index, cell in enumerate(cells):
            for paragraph in cell.paragraphs:
                # Абзацный отступ и выключка по ширине — для текста, а в ячейке
                # они превращают две строки в лесенку.
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.bold = index == 0


def _add_body(document, text: str) -> None:
    """Текст раздела абзацами — по одному на строку.

    Отданный целиком, он становится ОДНИМ абзацем с мягкими переносами внутри,
    а выключка по ширине растягивает каждую строку, оборванную таким переносом,
    во всю ширину полосы. Ровно это и происходит с разделом 4.6, где четыре
    блока рамки склеены через перевод строки.
    """
    for chunk in (line.strip() for line in (text or "").split("\n")):
        if chunk:
            document.add_paragraph(chunk)


def _add_section(document, section: TZSection, level: int) -> None:
    document.add_heading(f"{section.number}. {section.title}", level=level)
    if section.body:
        _add_body(document, section.body)
    if section.rows:
        _add_rows(document, section.rows)
    for sub in section.subsections:
        _add_section(document, sub, min(level + 1, 4))


def _title_page(document, spec: TechnicalSpec) -> None:
    """Титул: гриф утверждения, наименование, тема, ссылка на стандарт.

    Гриф — пустой бланк, а не утверждение: кто подписывает документ, система не
    знает и знать не может. Пустая строка под подпись — это форма; вписанная в
    неё фамилия была бы подлогом.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    right = WD_ALIGN_PARAGRAPH.RIGHT
    center = WD_ALIGN_PARAGRAPH.CENTER

    _plain(document, "УТВЕРЖДАЮ", align=right)
    if spec.customer and spec.customer != NOT_SET:
        _plain(document, spec.customer, align=right)
    _plain(document, "______________ /______________________/", align=right)
    _plain(document, "«___» ______________ 20___ г.", align=right)
    _plain(document)

    _plain(document, "ТЕХНИЧЕСКОЕ ЗАДАНИЕ", align=center, bold=True, size_pt=18)
    _plain(document, "на проведение научного исследования", align=center)
    _plain(document)
    _plain(document, spec.topic, align=center, bold=True)
    _plain(document)
    _plain(document, f"Составлено по {_STANDARD}", align=center, italic=True)
    _plain(document)

    _add_rows(document, [("Заказчик", spec.customer), ("Основание", spec.basis)])


def _signatures(document) -> None:
    """Строки для подписей — пустые, как в бланке."""
    _plain(document)
    for role in ("Руководитель работы", "Исполнитель"):
        _plain(document, f"{role}  ______________________  "
                         "/______________________/")


def build_tz_docx(spec: TechnicalSpec) -> bytes:
    """Собрать .docx и вернуть его байтами.

    Байтами, а не файлом: вызывающий решает, писать на диск, отдавать по HTTP
    или и то и другое, и ему не приходится чистить временный файл, если запись
    не удалась.
    """
    from docx import Document

    document = Document()
    _typography(document)
    _title_page(document, spec)

    document.add_page_break()
    for section in spec.sections:
        if _visible(section):
            _add_section(document, section, level=1)

    missing = spec.unfilled()
    if missing:
        _plain(document)
        _plain(document,
               "Сведения не заданы для разделов: " + ", ".join(missing) + ". "
               "Они заполняются оператором в форме «Основание и приёмка» рамки "
               "исследования.",
               italic=True)
    _signatures(document)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def write_tz_files(spec: TechnicalSpec, directory, stamp: str
                   ) -> Tuple[Optional[str], Optional[str]]:
    """Записать оба файла и вернуть их имена, или None для того, что не вышло.

    Ничего не поднимает: документ, который не удалось сохранить, не должен
    останавливать исследование — оператор увидит в ленте, что файла нет, и это
    честнее упавшего прогона.
    """
    from pathlib import Path

    directory = Path(directory)
    docx_name = md_name = None
    try:
        directory.mkdir(parents=True, exist_ok=True)
        md_name = f"ТЗ_{stamp}.md"
        (directory / md_name).write_text(render_tz_markdown(spec), encoding="utf-8")
    except OSError:
        md_name = None
    try:
        docx_name = f"ТЗ_{stamp}.docx"
        (directory / docx_name).write_bytes(build_tz_docx(spec))
    except Exception:  # noqa: BLE001 — python-docx raises its own family
        docx_name = None
    return docx_name, md_name


__all__ = ["build_tz_docx", "render_tz_markdown", "write_tz_files"]
