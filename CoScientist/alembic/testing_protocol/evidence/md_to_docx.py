#!/usr/bin/env python3
"""Render the protocol markdown into a .docx matching the protocol's plain style."""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def add_code(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Pt(18)
        r = p.add_run(line or " ")
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x20, 0x20, 0x20)


def emit_runs(par, text):
    """Handle **bold**, *italic* and `code` inline."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def main(src: str, dst: str) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    lines = open(src, encoding="utf-8").read().splitlines()
    i, code = 0, None
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if code is None:
                code = []
            else:
                add_code(doc, code)
                doc.add_paragraph()
                code = None
            i += 1
            continue
        if code is not None:
            code.append(line)
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.strip().startswith("|") and i + 1 < len(lines) and set(
                lines[i + 1].replace("|", "").replace(":", "").strip()) <= {"-", " "}:
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if set("".join(cells).replace(":", "")) <= {"-", " "}:
                    i += 1
                    continue
                rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                for r, cells in enumerate(rows):
                    for c, text in enumerate(cells[:len(rows[0])]):
                        cell = t.cell(r, c)
                        cell.text = ""
                        par = cell.paragraphs[0]
                        emit_runs(par, text)
                        for run in par.runs:
                            run.font.size = Pt(10)
                            if r == 0:
                                run.bold = True
                doc.add_paragraph()
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif re.match(r"^\s*[-*] ", line):
            depth = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Pt(18 + 18 * depth)
            emit_runs(p, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            emit_runs(p, re.sub(r"^\s*\d+\. ", "", line))
        elif line.strip():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            emit_runs(p, line.strip())
        i += 1

    doc.save(dst)
    print(f"written {dst}")


main(sys.argv[1], sys.argv[2])
