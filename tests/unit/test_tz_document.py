"""Техническое задание по ГОСТ 19.201-78, собранное из рамки исследования.

Рамка заполняла постановку; заказчику нужен документ. Здесь проверяется то, за
что документ и отвечает: что он говорит ровно то, что подтвердил оператор, и
ничего сверх; что все восемь разделов стандарта на месте и в порядке п. 1.4; и
что готовый файл можно открыть — из чата по ссылке и из графа с карточки
«Постановка».

Run from the repo root:  pytest tests/unit/test_tz_document.py -q
"""
import pytest
from dotenv import load_dotenv

load_dotenv()

from CoScientist.context_init.models import ResearchFrame  # noqa: E402
from CoScientist.context_init.tz import (  # noqa: E402
    NOT_SET,
    apply_prose,
    spec_from_frame,
)
from CoScientist.context_init.tz_docx import write_tz_files  # noqa: E402
from CoScientist.graph.research.store import FRAME_ID, ResearchGraphStore  # noqa: E402

#: Разделы, которые стандарт перечисляет в п. 1.4, в его порядке.
GOST_SECTIONS = ("1", "2", "3", "4", "5", "6", "7", "8")


def _frame(**answers) -> ResearchFrame:
    """Рамка, в которой заполнены названные поля и только они."""
    frame = ResearchFrame.blank()
    for block in frame.blocks:
        for field in block.fields:
            if field.name in answers:
                field.value = answers[field.name]
                field.status = "уточнено оператором"
    return frame


# ── состав документа ────────────────────────────────────────────────────────

def test_the_document_has_every_section_the_standard_lists_in_order():
    """П. 1.4 перечисляет восемь разделов. Документ, в котором раздела нет,
    читатель не может отличить от документа, где раздел не требовался."""
    spec = spec_from_frame(ResearchFrame.blank())
    assert [s.number for s in spec.sections] == list(GOST_SECTIONS)


def test_the_requirements_section_keeps_the_standards_own_subsection_numbers():
    """4.1-4.6 соответствуют п. 2.4 стандарта. Номера сохранены, чтобы
    соответствие читалось без сверки."""
    spec = spec_from_frame(ResearchFrame.blank())
    fourth = spec.section("4")
    assert [s.number for s in fourth.subsections] == [
        "4.1", "4.2", "4.3", "4.4", "4.5", "4.6"]


def test_a_subsection_the_standard_has_and_research_has_not_is_named_not_dropped():
    """Маркировка, упаковка, транспортирование (2.4.6-2.4.8) к исследованию не
    применимы. Молча выброшенные, они сообщили бы, что их не требовалось;
    стандарт разрешает объединять разделы, и документ говорит, что объединил."""
    body = spec_from_frame(ResearchFrame.blank()).section("4").body
    assert "2.4.6-2.4.8" in body and "не" in body


# ── ничего не выдумано ──────────────────────────────────────────────────────

def test_an_empty_frame_produces_an_empty_document_not_a_plausible_one():
    """Придуманный заказчик попадёт в документ, который пойдёт людям, и будет
    читаться как факт. Пусто — честный результат опроса."""
    spec = spec_from_frame(ResearchFrame.blank())
    assert spec.customer == NOT_SET and spec.basis == NOT_SET
    # Всё, кроме четвёртого: у него есть собственный текст — оговорка о
    # разделах стандарта, которые к исследованию не применимы, — и подразделы,
    # и пусты как раз они.
    empty = set(spec.unfilled())
    assert empty >= {"1", "2", "3", "5", "6", "7", "8"}
    assert empty >= {"4.1", "4.2", "4.3", "4.4", "4.5", "4.6"}


def test_the_sections_the_new_questions_fill_are_exactly_the_ones_that_were_empty():
    """Блок «Основание и приёмка» добавлен ровно под разделы 2.2, 2.5а, 2.6 и
    2.7 — те, которые из дружеского запроса не выводятся."""
    before = spec_from_frame(ResearchFrame.blank()).unfilled()
    assert {"2", "5", "7"} <= set(before)

    after = spec_from_frame(_frame(
        basis_document="Договор № 17 от 03.02.2026",
        customer="ИТМО, НИИ наносистем",
        topic_name="Токсикологическое профилирование метаболитов борщевика",
        deliverables="Отчёт о НИР по ГОСТ 7.32, набор данных SMILES",
        stages="1. Сбор данных — март. 2. Профилирование — апрель.",
        acceptance="Приёмка комиссией НИИ по акту",
    ))
    assert {"2", "5", "7"}.isdisjoint(after.unfilled())
    assert after.customer == "ИТМО, НИИ наносистем"
    assert "Договор № 17" in after.section("2").rows[0][1]


def test_the_model_may_rewrite_a_section_but_never_fill_an_empty_one():
    """Гибрид держится на этом: факты берутся из подтверждённой рамки, модель
    даёт им связный вид. Иначе «Не задано» превращается в обязательство."""
    spec = spec_from_frame(_frame(formulation="Насколько токсичны метаболиты?"))
    assert spec.section("1").body != NOT_SET and spec.section("7").body == NOT_SET

    apply_prose(spec, {"1": "Работа посвящена оценке токсичности метаболитов.",
                       "7": "Работа выполняется в три этапа до конца года."})
    assert spec.section("1").body.startswith("Работа посвящена")
    assert spec.section("7").body == NOT_SET


def test_the_research_tasks_reach_the_document_as_tasks():
    """«OP-1» — машинный ключ плана, а не строка документа."""
    from CoScientist.context_init.models import FrameOperation

    frame = ResearchFrame.blank()
    frame.operations = [
        FrameOperation(operation_id="OP-1", statement="Собрать SMILES метаболитов"),
        FrameOperation(operation_id="OP-2", statement="Предсказать LD50"),
    ]
    rows = spec_from_frame(frame).section("4").subsections[0].rows
    assert [k for k, _ in rows] == ["Задача 1", "Задача 2"]
    assert rows[1][1] == "Предсказать LD50"


# ── переизложение ───────────────────────────────────────────────────────────

def _tasked(*statements) -> ResearchFrame:
    """Рамка с задачами исследования, сформулированными как в запросе."""
    from CoScientist.context_init.models import FrameOperation

    frame = _frame(formulation="Насколько токсичны метаболиты?")
    frame.operations = [
        FrameOperation(operation_id=f"OP-{i}", statement=s)
        for i, s in enumerate(statements, 1)]
    return frame


def test_every_section_with_content_is_offered_for_rewriting():
    """Раздел, собранный из таблицы, — тот же пересказ запроса, что и текст.
    Не отдав его модели, документ оставляют с сырыми строками рамки."""
    spec = spec_from_frame(_frame(
        formulation="Насколько токсичны метаболиты?",
        computational="RDKit, ADMETlab",      # 4.4 — только таблица
        datasets="PubChem, ChEMBL",           # 4.5 — только таблица
    ))
    offered = {p.number for p in spec.rewritable()}
    assert {"4.4", "4.5"} <= offered, offered
    assert "7" not in offered  # пустой раздел не предлагают

    apply_prose(spec, {"4.4": "Для проведения работ планируется использовать "
                              "вычислительные средства, приведённые в таблице."})
    assert spec.part("4.4").body.startswith("Для проведения работ")
    # Таблица при этом осталась таблицей.
    assert spec.part("4.4").rows


def test_the_standards_own_clause_is_not_handed_to_the_model():
    """Оговорка о п. 1.4 — утверждение документа о себе, а не пересказ рамки.
    Переписанная, она перестанет означать то, что означает."""
    spec = spec_from_frame(_frame(formulation="Вопрос"))
    assert spec.section("4").fixed
    assert "4" not in {p.number for p in spec.rewritable()}

    before = spec.section("4").body
    apply_prose(spec, {"4": "Требования к результату излагаются ниже."})
    assert spec.section("4").body == before


def test_a_task_is_replaced_by_its_number_not_by_its_place():
    """Модель возвращает задачи не все и не подряд. Задача, сместившаяся на
    строку, — это уже другая задача."""
    from CoScientist.context_init.tz import apply_tasks

    spec = spec_from_frame(_tasked("Собери литературные данные",
                                   "Проведи кластеризацию",
                                   "Предскажи LD50"))
    apply_tasks(spec, {3: "Предсказание значений LD50 для мыши",
                       1: "Сбор и систематизация литературных данных"})
    said = [v for _k, v in spec.part("4.1").rows]
    assert said == ["Сбор и систематизация литературных данных",
                    "Проведи кластеризацию",
                    "Предсказание значений LD50 для мыши"]


def test_a_task_the_model_did_not_rewrite_keeps_the_operators_wording():
    """Пустая формулировка и номер, которого нет, — это молчание модели, а не
    указание стереть строку."""
    from CoScientist.context_init.tz import apply_tasks

    spec = spec_from_frame(_tasked("Собери литературные данные"))
    apply_tasks(spec, {1: "   ", 7: "Задача, которой нет", "x": "мусор"})
    assert [v for _k, v in spec.part("4.1").rows] == ["Собери литературные данные"]


def test_the_theme_is_named_by_the_model_only_when_the_customer_did_not_name_it():
    """Название темы, данное заказчиком, — его формулировка; документ не вправе
    её переписывать."""
    from CoScientist.context_init.tz import apply_topic

    mine = spec_from_frame(_frame(formulation="Насколько токсичны метаболиты?"))
    assert not mine.topic_confirmed
    apply_topic(mine, "Оценка токсикологического профиля метаболитов")
    assert mine.topic == "Оценка токсикологического профиля метаболитов"

    theirs = spec_from_frame(_frame(topic_name="Профилирование метаболитов"))
    assert theirs.topic_confirmed
    apply_topic(theirs, "Совсем другая тема")
    assert theirs.topic == "Профилирование метаболитов"


def test_the_theme_taken_from_the_question_is_cut_between_words():
    """Титульный лист — первое, что видят. Обрыв посреди слова («предскажу LD50
    для мыш») выдаёт документ, собранный машиной, сильнее всего прочего."""
    long = ("Составлю полный токсикологический профиль метаболитов Heracleum "
            "sosnowskyi: соберу литературные данные о метаболитах и их SMILES, "
            "проведу кластеризацию по структурному сходству, предскажу LD50 "
            "для мыши для всех путей введения")
    topic = spec_from_frame(_frame(formulation=long)).topic
    assert topic.endswith("…") and len(topic) <= 161
    head = topic[:-1]
    assert long.startswith(head)          # это начало вопроса, а не пересказ
    assert long[len(head)] in " ,.;:"     # и оборвано оно на границе слова


def test_the_draft_offers_the_facts_and_nothing_to_invent():
    """То, что видит модель: заполненные разделы, задачи и таблицы. Пустых
    разделов в черновике нет — предложенный, он приглашает их заполнить."""
    from CoScientist.context_init.tz import prose_request

    spec = spec_from_frame(_tasked("Собери литературные данные"),
                           original_request="Автоматизируй профиль, пожалуйста")
    draft = prose_request(spec)
    assert "Собери литературные данные" in draft
    assert "Автоматизируй профиль" in draft          # контекст для формулировок
    assert "заказчиком не задано" in draft            # тему предлагает модель
    assert "7. Стадии и этапы разработки" not in draft  # пустой раздел не отдан


# ── файлы ───────────────────────────────────────────────────────────────────

def test_the_word_file_opens_and_carries_the_sections(tmp_path):
    docx = pytest.importorskip("docx")
    spec = spec_from_frame(_frame(
        formulation="Насколько токсичны метаболиты борщевика?",
        topic_name="Профилирование метаболитов",
        customer="ИТМО"))
    name, _ = write_tz_files(spec, tmp_path, "test")
    assert name and (tmp_path / name).is_file()

    document = docx.Document(str(tmp_path / name))
    headings = [p.text for p in document.paragraphs
                if p.style.name.startswith("Heading")]
    for number, title in ((s.number, s.title) for s in spec.sections):
        assert f"{number}. {title}" in headings
    # Значения, которые читатель ищет первыми, лежат в таблице, а не в прозе.
    said = "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert "ИТМО" in said


def test_both_formats_are_written_from_one_assembly(tmp_path):
    spec = spec_from_frame(_frame(formulation="Вопрос"))
    docx_name, md_name = write_tz_files(spec, tmp_path, "test")
    assert docx_name and md_name
    text = (tmp_path / md_name).read_text(encoding="utf-8")
    for number, title in ((s.number, s.title) for s in spec.sections):
        assert f"{number}. {title}" in text


def test_the_whole_document_is_black_and_of_one_typeface(tmp_path):
    """Шаблон Word красит заголовки в синий по теме «Office» и набирает их
    Calibri Light. В техническом задании это читается как веб-страница."""
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    spec = spec_from_frame(_frame(formulation="Вопрос", customer="ИТМО"))
    name, _ = write_tz_files(spec, tmp_path, "test")
    document = docx.Document(str(tmp_path / name))

    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        assert style.font.color.rgb == RGBColor(0x00, 0x00, 0x00), style_name
        assert style.font.name == "Times New Roman", style_name
        # Ссылка на тему документа сильнее явного значения: пока она на месте,
        # заголовок остаётся синим, сколько бы раз ему ни назначили чёрный.
        rpr = style.element.get_or_add_rPr()
        color, fonts = rpr.find(qn("w:color")), rpr.find(qn("w:rFonts"))
        assert color is None or color.get(qn("w:themeColor")) is None, style_name
        assert fonts is None or fonts.get(qn("w:asciiTheme")) is None, style_name

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            color = run.font.color
            assert color.type is None or color.rgb == RGBColor(0, 0, 0), run.text


def test_the_informal_request_is_kept_out_of_the_document_but_reaches_the_model(tmp_path):
    """«Автоматизируй, пожалуйста» — чужая реплика. Её место в рамке и в графе,
    где она и хранится; в ТЗ попадает то, что из неё сформулировано."""
    from CoScientist.context_init.tz import prose_request
    from CoScientist.context_init.tz_docx import render_tz_markdown

    ask = "Автоматизируй составление профиля, пожалуйста"
    spec = spec_from_frame(_frame(formulation="Насколько токсичны метаболиты?"),
                           original_request=ask)
    assert ask not in render_tz_markdown(spec)

    docx = pytest.importorskip("docx")
    name, _ = write_tz_files(spec, tmp_path, "test")
    document = docx.Document(str(tmp_path / name))
    said = "\n".join(p.text for p in document.paragraphs)
    said += "\n".join(c.text for t in document.tables
                      for r in t.rows for c in r.cells)
    assert ask not in said
    # А модель его видит — иначе формулировать разделы будет не из чего.
    assert ask in prose_request(spec)


def test_the_section_text_is_laid_out_as_paragraphs_not_as_one_block(tmp_path):
    """Выключка по ширине растягивает строку, оборванную мягким переносом, во
    всю ширину полосы. Раздел 4.6 склеен из четырёх блоков рамки через перевод
    строки — отданный одним абзацем, он печатается лесенкой с дырами."""
    pytest.importorskip("docx")
    from CoScientist.context_init.tz import apply_prose
    from CoScientist.context_init.tz_docx import build_tz_docx

    spec = spec_from_frame(_frame(formulation="Вопрос"))
    apply_prose(spec, {"1": "Первый абзац раздела.\n\nВторой абзац раздела."})
    body = build_tz_docx(spec)

    import io
    import zipfile

    document = zipfile.ZipFile(io.BytesIO(body)).read("word/document.xml").decode()
    assert "Первый абзац раздела." in document
    assert "Второй абзац раздела." in document
    assert "<w:br/>" not in document


def test_the_document_is_laid_out_for_a4(tmp_path):
    """Шаблон python-docx свёрстан под Letter — американский формат, который в
    российской организации вылезет за поля при печати."""
    docx = pytest.importorskip("docx")

    spec = spec_from_frame(_frame(formulation="Вопрос"))
    name, _ = write_tz_files(spec, tmp_path, "test")
    page = docx.Document(str(tmp_path / name)).sections[0]
    assert round(page.page_width.cm, 1) == 21.0
    assert round(page.page_height.cm, 1) == 29.7


def test_a_document_that_cannot_be_written_does_not_stop_the_research(tmp_path):
    """Исследование, остановленное тем, что не собрался .docx, хуже, чем
    исследование без .docx."""
    import CoScientist.context_init.tz_docx as mod

    spec = spec_from_frame(ResearchFrame.blank())

    def explode(_spec):
        raise RuntimeError("python-docx недоступен")

    original, mod.build_tz_docx = mod.build_tz_docx, explode
    try:
        docx_name, md_name = write_tz_files(spec, tmp_path, "test")
    finally:
        mod.build_tz_docx = original
    assert docx_name is None and md_name


# ── стадия ──────────────────────────────────────────────────────────────────

def test_the_models_restatement_reaches_the_document_it_was_written_for(
        tmp_path, monkeypatch):
    """Разобранный `output_key` ADK кладёт в state_delta финального события, а
    Runner применит его, только когда это событие до него дойдёт. Документ
    собирается раньше — внутри `_emit_final`, до `yield`, — и в состоянии ещё
    пусто. Читая одно лишь состояние, стадия выпускала бы документ без единой
    переформулировки: ровно ту раскладку запроса по разделам, ради исправления
    которой она и заведена."""
    import json
    from types import SimpleNamespace

    from CoScientist.context_init import tz_agent as mod

    store = ResearchGraphStore(directory=str(tmp_path / "graph"))
    store.ensure_root("Насколько токсичны метаболиты?")
    monkeypatch.setattr(mod, "get_research_graph", lambda _ctx: store)

    written = {}

    def fake_write(spec, directory, stamp):
        written["spec"] = spec
        return "ТЗ_test.docx", "ТЗ_test.md"

    monkeypatch.setattr(mod, "write_tz_files", fake_write)

    agent = mod.TZSpecSessionAgent(name="TZSpecAgent", output_key="tz_prose")
    ctx = SimpleNamespace(
        session=SimpleNamespace(state={
            "research_frame": _frame(
                formulation="Насколько токсичны метаболиты?").model_dump()}),
        invocation_id="inv-1", branch="main")

    # То, чем модель отвечает: JSON-строка в тексте финального события.
    said = json.dumps({
        "topic": "Оценка токсикологического профиля метаболитов",
        "sections": [{"number": "1",
                      "text": "Работа посвящена оценке токсичности метаболитов."}],
    }, ensure_ascii=False)

    events = list(agent._post_final_events(ctx, said))

    spec = written["spec"]
    assert spec.section("1").body == "Работа посвящена оценке токсичности метаболитов."
    assert spec.topic == "Оценка токсикологического профиля метаболитов"
    assert len(events) == 1
    assert "ТЗ_test.docx" in (events[0].content.parts[0].text or "")


def test_the_raw_json_of_a_reasoning_model_does_not_reach_the_chat():
    """ТЗ теперь пишет рассуждающая модель, а она кладёт мысль первой частью и
    ответ второй. Погасив только нулевую, лента показала бы пользователю сырой
    JSON рядом с объявлением о документе."""
    import asyncio

    from google.adk.events.event import Event
    from google.genai import types

    from CoScientist.hitl.session_agent import SessionAgent

    class Stub:
        def _post_final_events(self, ctx, output_text):
            yield "объявление о документе"

    event = Event(invocation_id="inv", author="TZSpecAgent",
                  content=types.Content(role="model", parts=[
                      types.Part(text="мысль модели"),
                      types.Part(text='{"sections": []}')]))

    async def run():
        return [e async for e in SessionAgent._emit_final(Stub(), None, event, "")]

    emitted = asyncio.run(run())
    assert emitted[0] is event and emitted[1] == "объявление о документе"
    assert all(not (p.text or "") for p in event.content.parts)


# ── граф ────────────────────────────────────────────────────────────────────

@pytest.fixture
def store(tmp_path):
    s = ResearchGraphStore(directory=str(tmp_path))
    s.ensure_root("Насколько токсичны метаболиты борщевика?")
    return s


def _frame_card(store):
    """Карточка «Постановка», или None — она проекция и рисуется, только когда
    есть что показать."""
    return next((n for n in store.to_view()["nodes"] if n["id"] == FRAME_ID), None)


def _attached_specs(card):
    return [a for a in (card or {}).get("attachments", []) if a["kind"] == "spec"]


def test_the_specification_hangs_off_the_framing_card(store):
    """Панель делает ссылку из вложения и инертный текст из атрибута, так что
    это единственный способ открыть документ из графа."""
    from CoScientist.context_init.tz_agent import publish_spec

    spec = spec_from_frame(_frame(formulation="Вопрос"))
    docx_name, _ = publish_spec(spec, store, store._dir / "tz")
    assert docx_name

    specs = _attached_specs(_frame_card(store))
    assert len(specs) == 1, specs
    assert specs[0]["href"].startswith("/api/tz-document?name=")


def test_a_specification_derived_from_a_conclusion_stays_with_it(store):
    """Условие, а не перенос типа: ТЗ пишется до исследования, а спецификация,
    выведенная из вывода, принадлежит выводу."""
    r = store.commit(source="OrchestratorAgent", enforce_permissions=False, nodes=[
        {"type": "Conclusion", "ref": "cl_new", "attrs": {"synthesis": "Вывод."}},
        {"type": "Spec", "ref": "sp_new", "attrs": {"content": "Спецификация."}},
    ], edges=[{"type": "derived_from", "from": "#sp_new", "to": "#cl_new"}])
    assert r.ok, r.errors
    assert not _attached_specs(_frame_card(store))
    # И лежит там, откуда выведена.
    conclusion = next(n for n in store.to_view()["nodes"] if n["kind"] == "conclusion")
    assert [a["kind"] for a in conclusion["attachments"]] == ["spec"]


def test_the_link_the_card_carries_is_the_one_the_chat_carries():
    """Один адрес в двух местах: `chat.js` делает ссылку из голого
    `/api/tz-document…`, `store._href` пропускает `/api/…` в карточку. Разойдись
    они — одна из двух ссылок молча перестанет открываться."""
    from CoScientist.context_init.tz_agent import _link, announcement

    said = announcement("ТЗ_20260923.docx", "ТЗ_20260923.md")
    assert _link("ТЗ_20260923.docx") in said
    assert "ГОСТ 19.201-78" in said
    # Тело документа в ленту не выгружается.
    assert len(said.splitlines()) <= 4
